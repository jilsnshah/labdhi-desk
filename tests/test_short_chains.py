"""Long, tangled chains of trades with shorts - every figure checked by hand.

Each step of a chain states what the book must say afterwards, worked out on
paper: how much of each sale is short, each sale's margin, and the stock in
each warehouse. After every step the raw rows are audited and every screen is
checked against every other (summary vs stock rows vs positions vs
warehouses vs the sales' own margins).

Rates are paise per kg; 1 MT = 1000 kg, so 1 MT at a Rs 5/kg margin = Rs 5,000
= 500000 paise.

    python -m tests.test_short_chains
"""
import io
import unittest

from tests.common import MT, buy, fresh, party, product, sell, warehouse
from tests.test_edit import audit
from backend import api, db                                        # noqa: E402
from backend.services import deals, report, revise, stock          # noqa: E402

RS = 100          # paise in a rupee


def rs_per_mt(margin_per_kg_rs, mt):
    """Margin in paise of `mt` MT at Rs `margin_per_kg_rs` per kg."""
    return int(round(margin_per_kg_rs * RS * mt * 1000))


class Chain(unittest.TestCase):
    maxDiff = None

    # ---------------------------------------------------------------- helpers
    def setUp(self):
        fresh()
        self.M, self.A = warehouse("Mundra"), warehouse("Aslali")
        self.P = product()
        self.Q = product(grade="SG5", manufacturer="Reliance")

    def cell(self, wh, pid=None):
        rows = api.list_stock(product_id=pid or self.P, warehouse_id=wh, limit=10)["items"]
        return rows[0]["stock_g"] if rows else 0

    def short(self, deal_id):
        return deals.get_deal(deal_id)["short_g"]

    def margin(self, deal_id):
        return deals.get_deal(deal_id)["margin_paise"]

    def lot(self, deal_id):
        return db.q1("SELECT id FROM lots WHERE deal_id=? AND parent_lot_id IS NULL", (deal_id,))["id"]

    def screens_agree(self):
        """Every screen is a sum over the same rows - they cannot disagree."""
        audit(self)
        s = api.summary()
        rows = api.list_stock(limit=200)["items"]
        self.assertEqual(s["stock_g"], sum(r["stock_g"] for r in rows), "summary vs stock rows")
        pos = api.list_positions(include_flat=True, limit=200)["items"]
        self.assertEqual(s["stock_g"], sum(p["stock_g"] for p in pos), "summary vs positions")
        for p in pos:
            self.assertEqual(p["stock_g"], sum(w["stock_g"] for w in p["warehouses"]), "position vs its warehouses")
            self.assertEqual(p["short_g"], sum(w["short_g"] for w in p["warehouses"]))
            detail = api.get_position(p["product_id"])
            self.assertEqual((detail["stock_g"], detail["short_g"]), (p["stock_g"], p["short_g"]))
        whs = api.list_warehouses(limit=200)["items"]
        self.assertEqual(s["stock_g"], sum(w["stock_g"] for w in whs), "summary vs warehouses")
        sales = [deals.get_deal(r["id"]) for r in db.q("SELECT id FROM deals WHERE side='sell' AND status='booked'")]
        self.assertEqual(s["realised_paise"], sum(d["margin_paise"] for d in sales), "realised vs sales")
        short_total = sum(d["short_g"] for d in sales)
        self.assertEqual(sum(r["short_g"] for r in rows), short_total, "shorts on stock screen vs sales")
        # per warehouse filter of positions agrees with the cells
        for w in whs:
            fp = api.list_positions(include_flat=True, warehouse_id=w["id"], limit=200)["items"]
            self.assertEqual(sum(p["stock_g"] for p in fp), w["stock_g"], "positions filtered to %s" % w["name"])
        # the Excel sauda report's margin total is the book's realised profit
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(report.tape_workbook(status="booked")))
        ws = wb["Saudas"]
        head = [c.value for c in ws[4]]
        mi, si = head.index("Margin ₹"), head.index("Sold short, not covered (kg)")
        margins = shorts = 0
        for row in ws.iter_rows(min_row=5, values_only=True):
            if row[1] and str(row[1]).startswith("LE/"):
                margins += row[mi] or 0
                shorts += row[si] or 0
        self.assertAlmostEqual(margins, s["realised_paise"] / 100, places=2)
        self.assertAlmostEqual(shorts, short_total / 1000, places=3)
        report.flow_workbook()                       # builds without error

    def step(self, expect):
        """expect: {"short": {id: g}, "margin": {id: paise}, "cell": {(wh, pid): g}}"""
        for did, g in expect.get("short", {}).items():
            self.assertEqual(self.short(did), g, "short of deal %d" % did)
        for did, p in expect.get("margin", {}).items():
            self.assertEqual(self.margin(did), p, "margin of deal %d" % did)
        for (wh, pid), g in expect.get("cell", {}).items():
            self.assertEqual(self.cell(wh, pid), g, "stock of product %s in warehouse %s" % (pid, wh))
        self.screens_agree()

    def edit_as_previewed(self, deal_id, changes, **consent):
        """The preview must say exactly what saving does."""
        p = revise.preview_edit(deal_id, changes, **consent)
        self.assertIsNone(p["error"], p["error"])
        want = {s["id"]: (s["margin_after_paise"], s["short_after_g"]) for s in p["effects"]["sales"]}
        before_realised = api.summary()["realised_paise"]
        revise.edit_deal(deal_id, changes, rehome=p["needs_rehome"], allow_short=p["needs_short"])
        for sid, (m, sh) in want.items():
            d = deals.get_deal(sid)
            if d["status"] == "booked":
                self.assertEqual((d["margin_paise"], d["short_g"]), (m, sh), "preview vs save for %s" % d["ref"])
        self.assertEqual(api.summary()["realised_paise"] - before_realised,
                         p["effects"]["realised_after_paise"] - p["effects"]["realised_before_paise"])
        return p

    def cancel_as_previewed(self, deal_id):
        p = revise.preview_cancel(deal_id)
        self.assertIsNone(p["error"], p["error"])
        want = {s["id"]: (s["margin_after_paise"], s["short_after_g"]) for s in p["effects"]["sales"]}
        deals.cancel_deal(deal_id, rehome=p["needs_rehome"])
        for sid, (m, sh) in want.items():
            d = deals.get_deal(sid)
            if d["status"] == "booked":
                self.assertEqual((d["margin_paise"], d["short_g"]), (m, sh))
        return p

    # ---------------------------------------------------------------- the train
    def test_the_train(self):
        M, A, P, Q = self.M, self.A, self.P, self.Q

        # 1-2. a purchase, a covered sale
        L1 = buy("S1", 10 * MT, 9500, "2026-09-01", wh="Mundra")["id"]
        X = sell("B1", 6 * MT, 10000, "2026-09-02", wh="Mundra")["id"]
        self.step({"margin": {X: rs_per_mt(5, 6)}, "cell": {(M, P): 4 * MT}})

        # 3. Y takes the last 4 and is 3 short
        with self.assertRaises(deals.DealError):
            sell("B2", 7 * MT, 10100, "2026-09-03", wh="Mundra")          # not asked for
        Y = sell("B2", 7 * MT, 10100, "2026-09-03", wh="Mundra", allow_short=True)["id"]
        self.step({"short": {Y: 3 * MT}, "margin": {Y: rs_per_mt(6, 4)}, "cell": {(M, P): -3 * MT}})

        # 4. Z is wholly short
        Z = sell("B3", 5 * MT, 10200, "2026-09-04", wh="Mundra", allow_short=True)["id"]
        self.step({"short": {Z: 5 * MT}, "margin": {Z: 0}, "cell": {(M, P): -8 * MT}})

        # 5. Y made more short - needs asking; the preview says so
        with self.assertRaises(deals.ShortNeeded):
            revise.edit_deal(Y, {"qty_g": 10 * MT})
        p = self.edit_as_previewed(Y, {"qty_g": 10 * MT})
        self.assertTrue(p["needs_short"])
        self.step({"short": {Y: 6 * MT, Z: 5 * MT}, "margin": {Y: rs_per_mt(6, 4)}, "cell": {(M, P): -11 * MT}})

        # 6. 8 MT arrive: Y (older) is covered in full first, then 2 of Z
        L2 = buy("S2", 8 * MT, 9600, "2026-09-05", wh="Mundra")["id"]
        self.step({"short": {Y: 0, Z: 3 * MT},
                   "margin": {Y: rs_per_mt(6, 4) + rs_per_mt(5, 6), Z: rs_per_mt(6, 2)},
                   "cell": {(M, P): -3 * MT}})

        # 7. stock in the other warehouse does not cover Mundra's short
        L3 = buy("S3", 5 * MT, 9700, "2026-09-06", wh="Aslali")["id"]
        self.step({"short": {Z: 3 * MT}, "cell": {(M, P): -3 * MT, (A, P): 5 * MT}})

        # 8. moving 4 MT of it to Mundra covers Z's 3 at that lot's cost
        mv = stock.transfer(self.lot(L3), M, 4 * MT)
        self.step({"short": {Z: 0}, "margin": {Z: rs_per_mt(6, 2) + rs_per_mt(5, 3)},
                   "cell": {(M, P): 1 * MT, (A, P): 1 * MT}})

        # 9. that transfer's stock is sold now: it cannot be undone
        with self.assertRaises(ValueError):
            stock.cancel_move(mv["id"])
        self.screens_agree()

        # 10. Z moves to Aslali: it hands its Mundra stock back and is short there
        with self.assertRaises(deals.ShortNeeded):
            revise.edit_deal(Z, {"warehouse_id": A})
        self.edit_as_previewed(Z, {"warehouse_id": A})
        self.step({"short": {Z: 4 * MT}, "margin": {Z: rs_per_mt(5, 1)},
                   "cell": {(M, P): 6 * MT, (A, P): -4 * MT}})

        # 11. 3 MT bought into Aslali cover 3 of Z's 4
        L4 = buy("S4", 3 * MT, 9400, "2026-09-07", wh="Aslali")["id"]
        self.step({"short": {Z: 1 * MT}, "margin": {Z: rs_per_mt(5, 1) + rs_per_mt(8, 3)},
                   "cell": {(A, P): -1 * MT}})

        # 12. that purchase really went to Mundra: Z is 4 short again, Mundra gains 3
        p = self.edit_as_previewed(L4, {"warehouse_id": M})
        self.assertTrue(p["needs_rehome"])
        self.step({"short": {Z: 4 * MT}, "margin": {Z: rs_per_mt(5, 1)},
                   "cell": {(M, P): 9 * MT, (A, P): -4 * MT}})

        # 13. X is cancelled: its 6 MT come back to Mundra
        self.cancel_as_previewed(X)
        self.step({"cell": {(M, P): 15 * MT, (A, P): -4 * MT}})

        # 14. Z comes back to Mundra and is covered by the oldest stock there (L1 at 95)
        self.edit_as_previewed(Z, {"warehouse_id": M})
        self.step({"short": {Z: 0}, "margin": {Z: rs_per_mt(7, 5)},
                   "cell": {(M, P): 10 * MT, (A, P): 1 * MT}})

        # 15. Undo: the last undoable action is L4's booking - it holds nothing now
        ev = deals.last_undoable()
        self.assertEqual((ev["entity"], ev["entity_id"]), ("deal", L4))
        deals.undo_event(int(ev["id"]))
        self.step({"cell": {(M, P): 7 * MT}})

        # 16. Undo again: the transfer - its stock is free again, so it goes back
        ev = deals.last_undoable()
        self.assertEqual((ev["entity"], ev["entity_id"]), ("move", mv["id"]))
        deals.undo_event(int(ev["id"]))
        self.step({"short": {Y: 0, Z: 0},
                   "margin": {Y: rs_per_mt(6, 4) + rs_per_mt(5, 6), Z: rs_per_mt(7, 5)},
                   "cell": {(M, P): 3 * MT, (A, P): 5 * MT}})

        # 17. L2 was really 5 MT: 1 MT of Y moves onto L1's last tonne
        self.edit_as_previewed(L2, {"qty_g": 5 * MT})
        self.step({"short": {Y: 0}, "margin": {Y: rs_per_mt(6, 5) + rs_per_mt(5, 5)}, "cell": {(M, P): 0}})

        # 18. ...no, 2 MT: nothing else is left in Mundra, so Y is 3 short again
        self.edit_as_previewed(L2, {"qty_g": 2 * MT})
        self.step({"short": {Y: 3 * MT}, "margin": {Y: rs_per_mt(6, 5) + rs_per_mt(5, 2)},
                   "cell": {(M, P): -3 * MT}})

        # 19. and at Rs 90: Y's 2 MT from it now earn Rs 11/kg
        self.edit_as_previewed(L2, {"rate_paise": 9000})
        self.step({"margin": {Y: rs_per_mt(6, 5) + rs_per_mt(11, 2)}})

        # 20. 4 MT at 99 cover Y's 3
        L5 = buy("S1", 4 * MT, 9900, "2026-09-08", wh="Mundra")["id"]
        self.step({"short": {Y: 0}, "margin": {Y: rs_per_mt(6, 5) + rs_per_mt(11, 2) + rs_per_mt(2, 3)},
                   "cell": {(M, P): 1 * MT}})

        # 21. Y renegotiated to Rs 98: 5 @95 +3, 2 @90 +8, 3 @99 -1
        self.edit_as_previewed(Y, {"rate_paise": 9800})
        self.step({"margin": {Y: rs_per_mt(3, 5) + rs_per_mt(8, 2) + rs_per_mt(-1, 3)}})

        # 22. L5 cancelled: Y is 3 short again (the 1 MT left of L5 goes with it)
        self.cancel_as_previewed(L5)
        self.step({"short": {Y: 3 * MT}, "margin": {Y: rs_per_mt(3, 5) + rs_per_mt(8, 2)},
                   "cell": {(M, P): -3 * MT}})

        # 22b. L2 was 4 MT after all: its 2 extra MT (at Rs 90) cover 2 of Y's 3 by editing alone
        self.edit_as_previewed(L2, {"qty_g": 4 * MT})
        self.step({"short": {Y: 1 * MT}, "margin": {Y: rs_per_mt(3, 5) + rs_per_mt(8, 4)},
                   "cell": {(M, P): -1 * MT}})

        # 23. a short sale in an empty warehouse, of another product, to the gram
        W = sell("B4", 2 * MT, 10000, "2026-09-09", wh="Aslali", allow_short=True,
                 grade="SG5", manufacturer="Reliance")["id"]
        self.step({"short": {W: 2 * MT}, "cell": {(A, Q): -2 * MT, (A, P): 5 * MT}})
        self.edit_as_previewed(W, {"qty_g": 1500})                      # 1.5 kg
        self.step({"short": {W: 1500}, "cell": {(A, Q): -1500}})
        buy("S2", 1 * MT, 9000, "2026-09-10", wh="Aslali", grade="SG5", manufacturer="Reliance")
        self.step({"short": {W: 0}, "margin": {W: 1500 * 1000 // 1000}, "cell": {(A, Q): 1 * MT - 1500}})

        # 24. W changes product to P in Aslali, which has 5 MT: fully covered, Q is back to 1 MT
        self.edit_as_previewed(W, {"product_id": P, "qty_g": 2 * MT})
        self.step({"short": {W: 0}, "margin": {W: rs_per_mt(3, 2)},
                   "cell": {(A, Q): 1 * MT, (A, P): 3 * MT}})

        # 25. everything cancelled, newest first: the book ends flat and exact
        for did in sorted((r["id"] for r in db.q("SELECT id FROM deals WHERE status='booked'")), reverse=True):
            p = revise.preview_cancel(did)
            deals.cancel_deal(did, rehome=p["needs_rehome"])
            self.screens_agree()
        s = api.summary()
        self.assertEqual((s["stock_g"], s["realised_paise"]), (0, 0))
        self.assertEqual(api.list_stock(limit=50)["items"], [])


class Weird(unittest.TestCase):
    """Odd corners, each on a small book."""

    def setUp(self):
        fresh()
        self.M, self.A, self.P = warehouse("Mundra"), warehouse("Aslali"), product()

    def check(self):
        Chain.screens_agree(self)

    cell = Chain.cell
    short = Chain.short
    margin = Chain.margin

    def test_two_shorts_the_same_day_are_covered_in_booking_order(self):
        a = sell("B1", 2 * MT, 10000, "2026-09-05", allow_short=True)["id"]
        b = sell("B2", 2 * MT, 10000, "2026-09-05", allow_short=True)["id"]
        buy("S1", 3 * MT, 9000, "2026-09-06")
        self.assertEqual((self.short(a), self.short(b)), (0, 1 * MT))
        self.check()

    def test_an_older_dated_short_is_covered_first_even_if_booked_later(self):
        late = sell("B1", 2 * MT, 10000, "2026-09-09", allow_short=True)["id"]
        early = sell("B2", 2 * MT, 10000, "2026-09-02", allow_short=True)["id"]
        buy("S1", 2 * MT, 9000, "2026-09-10")
        self.assertEqual((self.short(early), self.short(late)), (0, 2 * MT))
        self.check()

    def test_found_stock_covers_and_then_cannot_be_undone(self):
        b = buy("S1", 2 * MT, 9000, "2026-09-01")["id"]
        s = sell("B1", 5 * MT, 10000, "2026-09-02", allow_short=True)["id"]
        lot = db.q1("SELECT id FROM lots WHERE deal_id=?", (b,))["id"]
        mv = stock.adjust(lot, 1 * MT, reason="recount")
        self.assertEqual(self.short(s), 2 * MT)
        with self.assertRaises(ValueError):
            stock.cancel_move(mv["id"])
        self.check()

    def test_undoing_a_transfer_out_brings_stock_back_to_cover(self):
        b = buy("S1", 5 * MT, 9000, "2026-09-01")["id"]
        lot = db.q1("SELECT id FROM lots WHERE deal_id=?", (b,))["id"]
        mv = stock.transfer(lot, self.A, 3 * MT)
        s = sell("B1", 4 * MT, 10000, "2026-09-02", allow_short=True)["id"]      # 2 in Mundra, 2 short
        self.assertEqual(self.short(s), 2 * MT)
        stock.cancel_move(mv["id"])                                             # the 3 come home
        self.assertEqual(self.short(s), 0)
        self.assertEqual((self.cell(self.M), self.cell(self.A)), (1 * MT, 0))
        self.assertEqual(self.margin(s), 4 * MT // 1000 * 1000)
        self.check()

    def test_writing_off_leaves_no_free_stock_beside_a_short(self):
        b = buy("S1", 5 * MT, 9000, "2026-09-01")["id"]
        s = sell("B1", 3 * MT, 10000, "2026-09-02")["id"]
        lot = db.q1("SELECT id FROM lots WHERE deal_id=?", (b,))["id"]
        stock.adjust(lot, -2 * MT, reason="damaged")
        self.assertEqual(self.cell(self.M), 0)
        revise.edit_deal(s, {"qty_g": 4 * MT}, allow_short=True)
        self.assertEqual(self.short(s), 1 * MT)
        self.check()

    def test_a_short_sale_back_to_exactly_the_stock(self):
        buy("S1", 5 * MT, 9000, "2026-09-01")
        s = sell("B1", 8 * MT, 10000, "2026-09-02", allow_short=True)["id"]
        revise.edit_deal(s, {"qty_g": 5 * MT})
        self.assertEqual((self.short(s), self.cell(self.M)), (0, 0))
        self.check()

    def test_a_short_sale_to_a_product_nobody_holds(self):
        buy("S1", 5 * MT, 9000, "2026-09-01")
        s = sell("B1", 3 * MT, 10000, "2026-09-02")["id"]
        q = product(grade="SG5", manufacturer="Reliance")
        with self.assertRaises(deals.ShortNeeded):
            revise.edit_deal(s, {"product_id": q})
        revise.edit_deal(s, {"product_id": q}, allow_short=True)
        self.assertEqual((self.short(s), self.cell(self.M), self.cell(self.M, q)), (3 * MT, 5 * MT, -3 * MT))
        self.check()

    def test_the_mark_follows_a_short_sale(self):
        s = sell("B1", 3 * MT, 10000, "2026-09-02", allow_short=True)
        mark = db.q1("SELECT rate_paise, source FROM marks WHERE product_id=?", (self.P,))
        self.assertEqual((mark["rate_paise"], mark["source"]), (10000, "sale " + s["ref"]))
        buy("S1", 5 * MT, 9000, "2026-09-03")                       # a purchase never moves the mark
        self.assertEqual(db.q1("SELECT rate_paise FROM marks WHERE product_id=?", (self.P,))["rate_paise"], 10000)
        # open P&L: the 3 covered MT are sold; 2 MT left at mark 100 over cost 90
        self.assertEqual(api.summary()["unrealised_paise"], 2 * 1000 * 1000)
        self.check()

    def test_setting_allows_shorts_without_asking(self):
        with db.tx() as conn:
            db.set_setting(conn, "allow_short_sales", "1")
        s = sell("B1", 3 * MT, 10000, "2026-09-02")["id"]
        self.assertEqual(self.short(s), 3 * MT)
        revise.edit_deal(s, {"qty_g": 4 * MT})
        self.assertEqual(self.short(s), 4 * MT)
        self.check()

    def test_reallocating_a_short_sale(self):
        buy("S1", 2 * MT, 9000, "2026-09-01")
        b = buy("S2", 2 * MT, 9500, "2026-09-01")["id"]
        s = sell("B1", 6 * MT, 10000, "2026-09-02", allow_short=True)["id"]
        lot = db.q1("SELECT id FROM lots WHERE deal_id=?", (b,))["id"]
        # picking only one lot would leave the other unused - refused
        with self.assertRaises(deals.DealError):
            deals.reallocate(s, pins=[{"lot_id": lot, "qty_g": 2 * MT}, {"lot_id": lot - 1, "qty_g": 1 * MT}])
        self.assertEqual(self.short(s), 2 * MT)
        self.check()

    def test_preview_of_a_short_edit_saves_nothing(self):
        buy("S1", 2 * MT, 9000, "2026-09-01")
        s = sell("B1", 2 * MT, 10000, "2026-09-02")["id"]
        before = (api.summary(), api.list_stock(limit=50)["items"])
        p = revise.preview_edit(s, {"qty_g": 9 * MT})
        self.assertTrue(p["needs_short"])
        self.assertEqual(p["short_grams"], 7 * MT)
        self.assertEqual(before, (api.summary(), api.list_stock(limit=50)["items"]))
        self.assertEqual(db.scalar("SELECT COUNT(*) FROM events WHERE action IN ('edit','cover')"), 0)

    def test_undo_toast_after_a_covering_buy(self):
        s = sell("B1", 3 * MT, 10000, "2026-09-02", allow_short=True)["id"]
        buy("S1", 5 * MT, 9000, "2026-09-03")
        self.assertEqual(self.short(s), 0)
        api.undo()                                                   # what the toast's Undo calls
        self.assertEqual((self.short(s), self.cell(self.M)), (3 * MT, -3 * MT))
        api.undo()                                                   # and then the sale itself
        self.assertEqual(self.cell(self.M), 0)
        self.check()


    def test_a_backdated_purchase_covers_too(self):
        s = sell("B1", 3 * MT, 10000, "2026-09-10", allow_short=True)["id"]
        buy("S1", 2 * MT, 9000, "2026-09-01")                         # dated before the sale
        self.assertEqual((self.short(s), self.margin(s)), (1 * MT, 2 * 1000 * 1000))
        self.check()

    def test_a_cover_across_two_costs_then_both_repriced(self):
        s = sell("B1", 5 * MT, 10000, "2026-09-02", allow_short=True)["id"]
        a = buy("S1", 2 * MT, 9000, "2026-09-03")["id"]
        b = buy("S2", 2 * MT, 9500, "2026-09-04")["id"]
        self.assertEqual((self.short(s), self.margin(s)), (1 * MT, 2 * 1000 * 1000 + 2 * 1000 * 500))
        revise.edit_deal(a, {"rate_paise": 9100})
        revise.edit_deal(b, {"rate_paise": 9400})
        self.assertEqual(self.margin(s), 2 * 1000 * 900 + 2 * 1000 * 600)
        self.check()

    def test_moving_a_covered_short_frees_stock_for_the_next_short(self):
        s1 = sell("B1", 3 * MT, 10000, "2026-09-02", allow_short=True)["id"]
        s2 = sell("B2", 2 * MT, 10100, "2026-09-03", allow_short=True)["id"]
        buy("S1", 3 * MT, 9000, "2026-09-04")                          # covers s1 only
        self.assertEqual((self.short(s1), self.short(s2)), (0, 2 * MT))
        revise.edit_deal(s1, {"warehouse_id": self.A}, allow_short=True)   # s1 leaves Mundra...
        self.assertEqual((self.short(s1), self.short(s2)), (3 * MT, 0))    # ...its 3 MT go to s2
        self.assertEqual((self.cell(self.M), self.cell(self.A)), (1 * MT, -3 * MT))
        self.check()

    def test_a_purchase_moved_away_takes_its_covers_with_it(self):
        s = sell("B1", 3 * MT, 10000, "2026-09-02", allow_short=True)["id"]
        b = buy("S1", 5 * MT, 9000, "2026-09-03")["id"]
        t = sell("B2", 1 * MT, 10000, "2026-09-02", wh="Aslali", allow_short=True)["id"]
        with self.assertRaises(deals.RehomeNeeded):
            revise.edit_deal(b, {"warehouse_id": self.A})
        revise.edit_deal(b, {"warehouse_id": self.A}, rehome=True)
        self.assertEqual((self.short(s), self.short(t)), (3 * MT, 0))
        self.assertEqual((self.cell(self.M), self.cell(self.A)), (-3 * MT, 4 * MT))
        self.check()

    def test_undo_of_a_buy_whose_cover_was_re_split_asks_instead(self):
        s = sell("B1", 3 * MT, 10000, "2026-09-02", allow_short=True)["id"]
        b = buy("S1", 5 * MT, 9000, "2026-09-03")["id"]
        lot = db.q1("SELECT id FROM lots WHERE deal_id=?", (b,))["id"]
        revise.edit_deal(s, {"pins": [{"lot_id": lot, "qty_g": 3 * MT}]})   # the trader re-picked it
        with self.assertRaises(deals.RehomeNeeded):
            deals.undo_event(int(deals.last_undoable()["id"]))
        self.check()

    def test_kilo_precision_keeps_every_total_exact(self):
        # 1.234 MT at odd rates: all values are whole paise, so totals equal their parts
        buy("S1", 1234 * 1000, 9537, "2026-09-01")
        s = sell("B1", 2345 * 1000, 10013, "2026-09-02", allow_short=True)["id"]
        buy("S2", 777 * 1000, 9611, "2026-09-03")
        self.assertEqual(self.short(s), (2345 - 1234 - 777) * 1000)
        self.assertEqual(self.margin(s), 1234 * (10013 - 9537) + 777 * (10013 - 9611))
        self.check()


class Fuzz(unittest.TestCase):
    """Random books: shorts, re-splits, transfers, finds, write-offs, undos and backdating."""

    screens_agree = Chain.screens_agree
    cell = Chain.cell

    def test_fuzz(self):
        import random
        for seed in range(10):
            with self.subTest(seed=seed):
                self.run_seed(random.Random(100 + seed))

    def run_seed(self, rng):
        fresh()
        whs = [warehouse("Mundra"), warehouse("Aslali")]
        prods = [product(), product(grade="SG5", manufacturer="Reliance")]
        day = lambda: "2026-09-%02d" % rng.randint(1, 28)                     # noqa: E731
        kg = lambda lo, hi: rng.randint(lo, hi) * 1000                          # noqa: E731
        done = 0
        for step in range(160):
            live = db.dicts(db.q("SELECT * FROM deals WHERE status='booked'"))
            sells = [d for d in live if d["side"] == "sell"]
            r = rng.random()
            try:
                if r < 0.22:
                    pid, wid = rng.choice(prods), rng.choice(whs)
                    deals.create_deal(dict(side="sell", party_id=party("B%d" % rng.randint(1, 4)), product_id=pid,
                                           warehouse_id=wid, qty_g=kg(100, 6000), allow_short=True,
                                           rate_paise=rng.randint(9000, 11000), deal_date=day()))
                elif r < 0.40:
                    deals.create_deal(dict(side="buy", party_id=party("S%d" % rng.randint(1, 4)),
                                           product_id=rng.choice(prods), warehouse_id=rng.choice(whs),
                                           qty_g=kg(100, 6000), rate_paise=rng.randint(9000, 10000), deal_date=day()))
                elif r < 0.47:
                    lots = db.dicts(db.q("SELECT id, warehouse_id FROM lots WHERE status='open'"))
                    if lots:
                        lot = rng.choice(lots)
                        free = stock.get_lot(lot["id"])["available_g"]
                        dest = [w for w in whs if w != lot["warehouse_id"]][0]
                        stock.transfer(lot["id"], dest, max(1000, (rng.randint(1, free) // 1000) * 1000))
                elif r < 0.52:
                    lots = db.dicts(db.q("SELECT id FROM lots WHERE status IN ('open','exhausted')"))
                    if lots:
                        amount = kg(1, 800) * rng.choice([1, -1])
                        stock.adjust(rng.choice(lots)["id"], amount)
                elif r < 0.60:
                    ev = deals.last_undoable()
                    if ev:
                        deals.undo_event(int(ev["id"]))
                elif r < 0.68 and sells:
                    d = rng.choice(sells)
                    lots = revise.lots_for_sale(d["product_id"], d["warehouse_id"], d["id"])
                    pins, left = [], d["qty_g"]
                    for l in rng.sample(lots, len(lots)):
                        take = min(l["available_g"], (rng.randint(0, left) // 1000) * 1000)
                        pins.append({"lot_id": l["id"], "qty_g": take}); left -= take
                    revise.edit_deal(d["id"], {"pins": pins}, allow_short=rng.random() < 0.7)
                elif live:
                    d = rng.choice(live)
                    field = rng.choice(["qty_g", "rate_paise", "warehouse_id", "product_id", "deal_date", "cancel"])
                    if field == "cancel":
                        deals.cancel_deal(d["id"], rehome=rng.random() < 0.8)
                    else:
                        value = {"qty_g": max(1000, d["qty_g"] + kg(-3000, 3000)), "rate_paise": rng.randint(8500, 11500),
                                 "warehouse_id": rng.choice(whs), "product_id": rng.choice(prods), "deal_date": day()}[field]
                        revise.edit_deal(d["id"], {field: value}, rehome=rng.random() < 0.8,
                                         allow_short=rng.random() < 0.8)
                done += 1
            except (deals.DealError, ValueError):
                pass
            if step % 20 == 19:
                self.screens_agree()
            else:
                audit(self)
        self.screens_agree()
        self.assertGreater(done, 60)


if __name__ == "__main__":
    unittest.main(verbosity=2)
