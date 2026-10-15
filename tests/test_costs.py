"""Labour and freight: ours on both sides, straight off the profit.

They are recorded on the sauda as two whole-deal amounts and counted on its
own date - a purchase's costs reduce profit the day it is bought, and stock
stays valued at the purchase rate. So every figure that says "profit" is
margin less costs, and nothing about stock, lots or allocations moves.

    python -m tests.test_costs
"""
import io
import unittest

from tests.common import MT, buy, conserved, fresh, party, product, sell, warehouse
from tests.test_edit import audit, book
from backend import api, db                                        # noqa: E402
from backend.services import deals, report, revise, stock          # noqa: E402

RS = 100           # paise in a rupee


class Costs(unittest.TestCase):
    maxDiff = None

    def setUp(self):
        fresh()
        self.M, self.A, self.P = warehouse("Mundra"), warehouse("Aslali"), product()

    def deal(self, deal_id):
        return deals.get_deal(deal_id)

    # ------------------------------------------------------------ booking
    def test_a_sale_earns_its_margin_less_its_costs(self):
        buy("S1", 10 * MT, 9500, "2026-09-01")
        s = sell("B1", 4 * MT, 10000, "2026-09-02",
                 labour_paise=2000 * RS, freight_paise=8000 * RS)["id"]
        d = self.deal(s)
        self.assertEqual(d["margin_paise"], 4000 * 500)              # 4 MT at Rs 5/kg = Rs 20,000
        self.assertEqual(d["costs_paise"], 10000 * RS)
        self.assertEqual(d["net_margin_paise"], 20000 * RS - 10000 * RS)
        self.assertEqual(d["margin_rate_paise"], 10000 * RS * 1000 // (4 * MT))   # Rs 2.50/kg
        audit(self)

    def test_a_purchase_costs_are_counted_the_day_it_is_bought(self):
        b = buy("S1", 10 * MT, 9500, "2026-09-01",
                labour_paise=5000 * RS, freight_paise=25000 * RS)["id"]
        self.assertEqual(self.deal(b)["costs_paise"], 30000 * RS)
        self.assertEqual(self.deal(b)["net_margin_paise"], -30000 * RS)
        # the stock itself is untouched: still 10 MT at Rs 95
        lot = db.q1("SELECT qty_g, rate_paise FROM lots WHERE deal_id=?", (b,))
        self.assertEqual((lot["qty_g"], lot["rate_paise"]), (10 * MT, 9500))
        s = api.summary()
        self.assertEqual(s["realised_paise"], -30000 * RS)
        self.assertEqual(s["stock_value_paise"], 10000 * 9500)       # 10 MT at Rs 95, no costs added
        audit(self)

    def test_the_book_adds_up_across_both_sides(self):
        buy("S1", 10 * MT, 9500, "2026-09-01", labour_paise=5000 * RS, freight_paise=25000 * RS)
        sell("B1", 4 * MT, 10000, "2026-09-02", labour_paise=2000 * RS, freight_paise=8000 * RS)
        s = api.summary()
        self.assertEqual(s["costs_paise"], 40000 * RS)
        self.assertEqual(s["realised_paise"], 20000 * RS - 40000 * RS)   # -Rs 20,000
        self.assertEqual(s["realised_today_paise"], 0)                   # both dated earlier
        self.assertEqual(api.get_position(self.P)["realised_paise"], -20000 * RS)
        who = {c["name"]: c["margin_paise"] for c in api.counterparties(limit=10)["items"]}
        self.assertEqual(who, {"B1": 10000 * RS, "S1": -30000 * RS})
        audit(self)

    def test_costs_are_counted_in_the_period_of_their_own_sauda(self):
        from backend.services import dashboard
        buy("S1", 10 * MT, 9500, "2026-09-01", freight_paise=25000 * RS)
        sell("B1", 4 * MT, 10000, "2026-09-20", freight_paise=8000 * RS)
        self.assertEqual(dashboard.realised_between("2026-09-01", "2026-09-05"), -25000 * RS)
        self.assertEqual(dashboard.realised_between("2026-09-06", "2026-09-30"), 20000 * RS - 8000 * RS)
        self.assertEqual(dashboard.realised_between("2026-09-01", "2026-09-30"), -13000 * RS)

    def test_a_negative_amount_is_refused(self):
        with self.assertRaises(deals.DealError):
            buy("S1", 1 * MT, 9500, "2026-09-01", labour_paise=-100)

    def test_costs_do_not_touch_stock_or_open_pnl(self):
        before = book()
        buy("S1", 10 * MT, 9500, "2026-09-01")
        plain = book()
        fresh()
        warehouse("Mundra"), warehouse("Aslali"), product()
        buy("S1", 10 * MT, 9500, "2026-09-01", labour_paise=5000 * RS, freight_paise=25000 * RS)
        withcosts = book()
        self.assertEqual(plain["stock"], withcosts["stock"])
        self.assertEqual(plain["positions"][0]["stock_g"], withcosts["positions"][0]["stock_g"])
        self.assertEqual(plain["positions"][0]["unrealised_paise"], withcosts["positions"][0]["unrealised_paise"])
        self.assertEqual(plain["lots"], withcosts["lots"])
        self.assertNotEqual(before, withcosts)

    # ------------------------------------------------------------ changing them
    def test_costs_can_be_edited_and_the_preview_says_what_changes(self):
        buy("S1", 10 * MT, 9500, "2026-09-01")
        s = sell("B1", 4 * MT, 10000, "2026-09-02", labour_paise=2000 * RS)["id"]
        p = revise.preview_edit(s, {"labour_paise": 3000 * RS, "freight_paise": 9000 * RS})
        self.assertIsNone(p["error"])
        self.assertEqual({c["label"]: (c["before"], c["after"]) for c in p["changes"]},
                         {"Labour": ("₹2,000", "₹3,000"), "Freight": ("₹0", "₹9,000")})
        effect = p["effects"]["sales"][0]
        self.assertEqual((effect["margin_before_paise"], effect["margin_after_paise"]),
                         (18000 * RS, 8000 * RS))
        revise.edit_deal(s, {"labour_paise": 3000 * RS, "freight_paise": 9000 * RS})
        self.assertEqual(self.deal(s)["net_margin_paise"], 8000 * RS)
        self.assertEqual(api.summary()["realised_paise"], 8000 * RS)
        audit(self)

    def test_editing_the_costs_moves_nothing_else(self):
        buy("S1", 10 * MT, 9500, "2026-09-01")
        s = sell("B1", 4 * MT, 10000, "2026-09-02")["id"]
        before = book()
        revise.edit_deal(s, {"freight_paise": 7000 * RS})
        after = book()
        for key in ("stock", "lots", "allocations", "marks", "warehouses"):
            self.assertEqual(before[key], after[key], key)
        self.assertEqual(api.summary()["realised_paise"], 20000 * RS - 7000 * RS)
        audit(self)

    def test_cancelling_takes_its_costs_off_the_book(self):
        buy("S1", 10 * MT, 9500, "2026-09-01", freight_paise=25000 * RS)
        s = sell("B1", 4 * MT, 10000, "2026-09-02", freight_paise=8000 * RS)["id"]
        self.assertEqual(api.summary()["realised_paise"], 20000 * RS - 33000 * RS)
        deals.cancel_deal(s)
        self.assertEqual(api.summary()["realised_paise"], -25000 * RS)    # only the purchase's freight
        self.assertEqual(self.deal(s)["costs_paise"], 0)                  # a cancelled sauda costs nothing
        audit(self)

    def test_undo_takes_them_off_too(self):
        buy("S1", 10 * MT, 9500, "2026-09-01", freight_paise=25000 * RS)
        before = api.summary()["realised_paise"]
        sell("B1", 4 * MT, 10000, "2026-09-02", labour_paise=1000 * RS)
        deals.undo_event(int(deals.last_undoable()["id"]))
        self.assertEqual(api.summary()["realised_paise"], before)
        audit(self)

    def test_a_short_sale_carries_its_costs_too(self):
        s = sell("B1", 3 * MT, 10000, "2026-09-02", allow_short=True, freight_paise=6000 * RS)["id"]
        self.assertEqual((self.deal(s)["margin_paise"], self.deal(s)["net_margin_paise"]), (0, -6000 * RS))
        buy("S1", 3 * MT, 9000, "2026-09-03", freight_paise=3000 * RS)     # covers it
        self.assertEqual(self.deal(s)["net_margin_paise"], 30000 * RS - 6000 * RS)
        self.assertEqual(api.summary()["realised_paise"], 30000 * RS - 9000 * RS)
        audit(self)

    # ------------------------------------------------------------ the reports
    def test_the_excel_report_shows_and_totals_them(self):
        from openpyxl import load_workbook
        buy("S1", 10 * MT, 9500, "2026-09-01", labour_paise=5000 * RS, freight_paise=25000 * RS)
        sell("B1", 4 * MT, 10000, "2026-09-02", labour_paise=2000 * RS, freight_paise=8000 * RS)
        ws = load_workbook(io.BytesIO(report.tape_workbook()))["Saudas"]
        head = [c.value for c in ws[4]]
        cols = {name: head.index(name) for name in ("Sauda No.", "Labour ₹", "Freight ₹", "Net after costs ₹")}
        rows = {r[cols["Sauda No."]]: r for r in ws.iter_rows(min_row=5, values_only=True)
                if r[1] and str(r[1]).startswith("LE/")}
        self.assertEqual([rows["LE/26-27/0001"][cols[c]] for c in ("Labour ₹", "Freight ₹", "Net after costs ₹")],
                         [5000, 25000, -30000])
        self.assertEqual([rows["LE/26-27/0002"][cols[c]] for c in ("Labour ₹", "Freight ₹", "Net after costs ₹")],
                         [2000, 8000, 10000])
        net = sum(rows[r][cols["Net after costs ₹"]] for r in rows)
        self.assertEqual(net * RS, api.summary()["realised_paise"])
        report.flow_workbook()

    def test_every_screen_agrees_after_costs(self):
        buy("S1", 10 * MT, 9500, "2026-09-01", labour_paise=5000 * RS, freight_paise=25000 * RS)
        buy("S2", 6 * MT, 9700, "2026-09-02", wh="Aslali", freight_paise=4000 * RS)
        sell("B1", 4 * MT, 10000, "2026-09-03", labour_paise=2000 * RS)
        sell("B2", 5 * MT, 10200, "2026-09-04", wh="Aslali", freight_paise=3000 * RS)
        s = api.summary()
        sales = [deals.get_deal(r["id"]) for r in db.q("SELECT id FROM deals WHERE side='sell' AND status='booked'")]
        buys = [deals.get_deal(r["id"]) for r in db.q("SELECT id FROM deals WHERE side='buy' AND status='booked'")]
        self.assertEqual(s["realised_paise"],
                         sum(d["net_margin_paise"] for d in sales) - sum(d["costs_paise"] for d in buys))
        self.assertEqual(s["costs_paise"], sum(d["costs_paise"] for d in sales + buys))
        self.assertEqual(s["realised_paise"],
                         sum(api.get_position(p["id"])["realised_paise"]
                             for p in api.list_products(limit=50)["items"]))
        conserved(self, self.P)
        audit(self)


if __name__ == "__main__":
    unittest.main(verbosity=2)
