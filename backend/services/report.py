"""The Flow of Material report, as an Excel workbook.

The same window the Flow screen draws: sales made in the period, every purchase
that fed them whatever its own date, and purchases made in the period even if
nothing of them is sold yet. Everything is read from the book at the moment of
export; nothing is stored.

Sheets:
    Summary          the period at a glance, and one line per product
    Material Flow    one row per allocation: which purchase, from whom, at what
                     cost, went to which sale, to whom, at what rate, for what margin
    Sales            one row per sale, with where its material came from
    Purchases        one row per purchase, with where its material went and what is left
    Stock Movements  transfers and adjustments in the period

Money is written as rupees and quantities as kg or MT, formatted as numbers so
the sheet can be summed, filtered and pivoted.
"""
from __future__ import annotations

import io
from collections import defaultdict
from datetime import date, datetime
from typing import Any, Dict, List, Optional

from .. import db

# ------------------------------------------------------------------ styling
INK, MUTED, LINE = "16181C", "5B6170", "D9DBD4"
BLUE, BLUE_SOFT = "2F6DF6", "EAF0FF"
GREEN, GREEN_SOFT, RED, RED_SOFT = "0D7A48", "E6F5ED", "C33636", "FDECEC"
HEAD_FILL = "1F2937"

KG = '#,##0'
MT = '#,##0.000'
RATE = '"₹"#,##0.00'
MONEY = '"₹"#,##,##0;[Red]-"₹"#,##,##0'
MARGIN_RATE = '"₹"#,##0.00;[Red]-"₹"#,##0.00'
PCT = '0.00%;[Red]-0.00%'
DATE = 'DD-MMM-YYYY'
DAYS = '0'


def _d(iso: Optional[str]):
    if not iso:
        return None
    try:
        return datetime.strptime(iso[:10], "%Y-%m-%d").date()
    except ValueError:
        return None


def _r(paise) -> float:
    return round((paise or 0) / 100.0, 2)


def _kg(g) -> float:
    return round((g or 0) / 1000.0, 3)


def _mt(g) -> float:
    return round((g or 0) / 1e6, 6)


def _money(g, paise_per_kg) -> float:
    """grams x paise/kg -> rupees, exact to the paisa."""
    return round((g or 0) * (paise_per_kg or 0) / 1000.0 / 100.0, 2)


# ------------------------------------------------------------------ data
def _window(date_from, date_to, product_id, warehouse_id):
    sale_where, sale_args = ["d.side = 'sell'", "d.status = 'booked'"], []
    buy_where, buy_args = ["d.side = 'buy'", "d.status = 'booked'"], []
    for where, args in ((sale_where, sale_args), (buy_where, buy_args)):
        if date_from:
            where.append("d.deal_date >= ?"); args.append(date_from)
        if date_to:
            where.append("d.deal_date <= ?"); args.append(date_to)
        if product_id:
            where.append("d.product_id = ?"); args.append(int(product_id))
        if warehouse_id:
            where.append("d.warehouse_id = ?"); args.append(int(warehouse_id))

    deal_cols = """d.*, p.name AS party_name, p.gstin AS party_gstin, s.display AS product,
                   s.material, s.grade, s.manufacturer, w.name AS warehouse, t.name AS transporter_name"""
    deal_from = """FROM deals d JOIN parties p ON p.id = d.party_id JOIN v_products s ON s.id = d.product_id
                   LEFT JOIN warehouses w ON w.id = d.warehouse_id LEFT JOIN parties t ON t.id = d.transporter_id"""
    sales = db.dicts(db.q("SELECT %s %s WHERE %s ORDER BY d.deal_date, d.id"
                          % (deal_cols, deal_from, " AND ".join(sale_where)), sale_args))

    flows: List[Dict[str, Any]] = []
    if sales:
        ids = [s["id"] for s in sales]
        flows = db.dicts(db.q(
            """SELECT a.id, a.sale_deal_id, a.lot_id, a.qty_g, a.cost_paise, a.sale_rate_paise,
                      l.deal_id AS buy_deal_id, lw.name AS lot_warehouse
               FROM allocations a JOIN lots l ON l.id = a.lot_id
               LEFT JOIN warehouses lw ON lw.id = l.warehouse_id
               WHERE a.active = 1 AND a.sale_deal_id IN (%s) ORDER BY a.sale_deal_id, a.id"""
            % ",".join("?" * len(ids)), ids))

    buy_ids = {r["id"] for r in db.q("SELECT d.id %s WHERE %s" % (deal_from, " AND ".join(buy_where)), buy_args)}
    buy_ids |= {f["buy_deal_id"] for f in flows}
    buys = []
    if buy_ids:
        ids = sorted(buy_ids)
        buys = db.dicts(db.q("SELECT %s %s WHERE d.id IN (%s) ORDER BY d.deal_date, d.id"
                             % (deal_cols, deal_from, ",".join("?" * len(ids))), ids))
    return sales, buys, flows


def _buy_outcomes(buy_ids: List[int]) -> Dict[int, Dict[str, Any]]:
    """For each purchase: everything sold out of it (ever), moved, written off, still here."""
    out = {i: {"sold_g": 0, "sold_value": 0.0, "margin": 0.0, "to": [], "moved_g": 0, "lost_g": 0,
               "found_g": 0, "left_g": 0, "left_where": defaultdict(int)} for i in buy_ids}
    if not buy_ids:
        return out
    marks = ",".join("?" * len(buy_ids))
    for r in db.q(
            """SELECT l.deal_id, a.qty_g, a.cost_paise, a.sale_rate_paise, sd.ref, sd.deal_date, cp.name AS buyer
               FROM allocations a JOIN lots l ON l.id = a.lot_id
               JOIN deals sd ON sd.id = a.sale_deal_id JOIN parties cp ON cp.id = sd.party_id
               WHERE a.active = 1 AND sd.status = 'booked' AND l.deal_id IN (%s)
               ORDER BY sd.deal_date, sd.id""" % marks, buy_ids):
        o = out[r["deal_id"]]
        o["sold_g"] += r["qty_g"]
        o["sold_value"] += _money(r["qty_g"], r["sale_rate_paise"])
        o["margin"] += _money(r["qty_g"], r["sale_rate_paise"] - r["cost_paise"])
        o["to"].append("%s %s MT (%s)" % (r["buyer"], _fmt_mt(r["qty_g"]), r["ref"]))
    for r in db.q(
            """SELECT l.deal_id, m.kind, m.qty_g FROM stock_moves m JOIN lots l ON l.id = m.lot_id
               WHERE m.status = 'done' AND l.deal_id IN (%s)""" % marks, buy_ids):
        o = out[r["deal_id"]]
        if r["kind"] == "transfer":
            o["moved_g"] += r["qty_g"]
        elif r["qty_g"] < 0:
            o["lost_g"] += -r["qty_g"]
        else:
            o["found_g"] += r["qty_g"]
    for r in db.q(
            """SELECT l.deal_id, w.name, SUM(l.qty_g - l.qty_allocated_g - l.qty_out_g) AS left_g
               FROM lots l JOIN warehouses w ON w.id = l.warehouse_id
               WHERE l.status = 'open' AND l.deal_id IN (%s) GROUP BY l.deal_id, w.name""" % marks, buy_ids):
        if r["left_g"]:
            out[r["deal_id"]]["left_g"] += r["left_g"]
            out[r["deal_id"]]["left_where"][r["name"]] += r["left_g"]
    return out


def _moves(date_from, date_to, product_id, warehouse_id):
    from .stock import _LEDGER
    where, args = ["mv.kind IN ('transfer_out','transfer_in','loss','gain')"], []
    for col, val in (("mv.date >=", date_from), ("mv.date <=", date_to)):
        if val:
            where.append(col + " ?"); args.append(val)
    for col, val in (("mv.product_id", product_id), ("mv.warehouse_id", warehouse_id)):
        if val:
            where.append(col + " = ?"); args.append(int(val))
    return db.dicts(db.q(
        """SELECT mv.*, s.display AS product, w.name AS warehouse, bd.ref AS buy_ref, sp.name AS supplier
           FROM (%s) mv JOIN v_products s ON s.id = mv.product_id JOIN warehouses w ON w.id = mv.warehouse_id
           JOIN lots l ON l.id = mv.lot_id JOIN deals bd ON bd.id = l.deal_id JOIN parties sp ON sp.id = l.supplier_id
           WHERE %s ORDER BY mv.date, mv.ts""" % (_LEDGER, " AND ".join(where)), args))


def _fmt_mt(g) -> str:
    return ("%.3f" % _mt(g)).rstrip("0").rstrip(".")


# ------------------------------------------------------------------ workbook
def flow_workbook(date_from: Optional[str] = None, date_to: Optional[str] = None,
                  product_id: Optional[int] = None, warehouse_id: Optional[int] = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    sales, buys, flows = _window(date_from, date_to, product_id, warehouse_id)
    sale_by_id = {s["id"]: s for s in sales}
    buy_by_id = {b["id"]: b for b in buys}
    outcomes = _buy_outcomes([b["id"] for b in buys])
    moves = _moves(date_from, date_to, product_id, warehouse_id)
    settings = db.settings()
    company = settings.get("company_name") or "Labdhi Exim"
    product_name = warehouse_name = None
    if product_id:
        product_name = db.scalar("SELECT display FROM v_products WHERE id=?", (int(product_id),), None)
    if warehouse_id:
        warehouse_name = db.scalar("SELECT name FROM warehouses WHERE id=?", (int(warehouse_id),), None)
    period = "%s to %s" % (_d(date_from).strftime("%d-%b-%Y") if date_from else "the beginning",
                           _d(date_to).strftime("%d-%b-%Y") if date_to else date.today().strftime("%d-%b-%Y"))

    wb = Workbook()
    thin = Side(style="thin", color=LINE)
    border = Border(bottom=thin)
    head_font = Font(bold=True, color="FFFFFF", size=10.5)
    head_fill = PatternFill("solid", fgColor=HEAD_FILL)
    total_fill = PatternFill("solid", fgColor=BLUE_SOFT)

    def table(ws, top: int, columns: List[tuple], rows: List[List[Any]], totals: Dict[str, Any] = None,
              zebra: bool = True, group_fill=None):
        """columns: (header, width, number_format or None). Writes a filterable,
        frozen table starting at row `top`; returns the row after it."""
        for c, (name, width, _fmt) in enumerate(columns, start=1):
            cell = ws.cell(row=top, column=c, value=name)
            cell.font, cell.fill = head_font, head_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(c)].width = width
        ws.row_dimensions[top].height = 32
        for i, row in enumerate(rows, start=1):
            fill = group_fill(i - 1) if group_fill else (PatternFill("solid", fgColor="F7F8F5") if zebra and i % 2 == 0 else None)
            for c, value in enumerate(row, start=1):
                cell = ws.cell(row=top + i, column=c, value=value)
                fmt = columns[c - 1][2]
                if fmt:
                    cell.number_format = fmt
                if fill:
                    cell.fill = fill
                cell.border = border
                cell.alignment = Alignment(vertical="top", wrap_text=isinstance(value, str) and len(value) > 38)
        last = top + len(rows)
        if rows:
            ws.auto_filter.ref = "A%d:%s%d" % (top, get_column_letter(len(columns)), last)
        ws.freeze_panes = ws.cell(row=top + 1, column=3)
        if totals and rows:
            t = last + 1
            ws.cell(row=t, column=1, value="TOTAL").font = Font(bold=True)
            for c, (name, _w, fmt) in enumerate(columns, start=1):
                if name in totals:
                    kind = totals[name]
                    col = get_column_letter(c)
                    if kind == "sum":
                        value = "=SUBTOTAL(9,%s%d:%s%d)" % (col, top + 1, col, last)
                    else:
                        value = kind(top + 1, last)
                    cell = ws.cell(row=t, column=c, value=value)
                    cell.number_format = fmt or KG
                    cell.font = Font(bold=True)
                ws.cell(row=t, column=c).fill = total_fill
            return t + 1
        return last + 1

    def title(ws, text, sub):
        ws["A1"] = "%s — %s" % (company.upper(), text)
        ws["A1"].font = Font(bold=True, size=15, color=INK)
        ws["A2"] = sub
        ws["A2"].font = Font(size=10.5, color=MUTED)
        ws.sheet_view.showGridLines = False

    scope = ["Period: %s" % period]
    if product_name:
        scope.append("Product: %s" % product_name)
    if warehouse_name:
        scope.append("Warehouse: %s" % warehouse_name)
    scope_text = "   ·   ".join(scope)

    # ================================================================ Material Flow
    ws = wb.active
    ws.title = "Material Flow"
    title(ws, "Flow of Material", scope_text + "   ·   one row per quantity that moved from a purchase to a sale")
    rows = []
    group = []
    for f in flows:
        s, b = sale_by_id[f["sale_deal_id"]], buy_by_id[f["buy_deal_id"]]
        margin_rate = f["sale_rate_paise"] - f["cost_paise"]
        held = (_d(s["deal_date"]) - _d(b["deal_date"])).days if s["deal_date"] and b["deal_date"] else None
        rows.append([
            _d(b["deal_date"]), b["ref"], b["party_name"], b["party_gstin"] or "",
            s["material"], s["grade"], s["manufacturer"], f["lot_warehouse"] or "",
            _kg(f["qty_g"]), _mt(f["qty_g"]),
            _r(f["cost_paise"]), _money(f["qty_g"], f["cost_paise"]),
            _d(s["deal_date"]), s["ref"], s["party_name"], s["party_gstin"] or "",
            _r(f["sale_rate_paise"]), _money(f["qty_g"], f["sale_rate_paise"]),
            _r(margin_rate), _money(f["qty_g"], margin_rate),
            (margin_rate / f["cost_paise"]) if f["cost_paise"] else None, held,
        ])
        group.append(s["id"])
    cols = [("Bought on", 12, DATE), ("Purchase Sauda No.", 16, None), ("Bought from (supplier)", 28, None),
            ("Supplier GSTIN", 17, None), ("Material", 10, None), ("Grade", 11, None), ("Manufacturer", 18, None),
            ("Warehouse", 12, None), ("Qty (kg)", 12, KG), ("Qty (MT)", 10, MT),
            ("Cost rate ₹/kg", 12, RATE), ("Cost value ₹", 14, MONEY),
            ("Sold on", 12, DATE), ("Sale Sauda No.", 16, None), ("Sold to (buyer)", 28, None),
            ("Buyer GSTIN", 17, None), ("Sale rate ₹/kg", 12, RATE), ("Sale value ₹", 14, MONEY),
            ("Margin ₹/kg", 12, MARGIN_RATE), ("Margin ₹", 14, MONEY), ("Margin %", 10, PCT), ("Days held", 9, DAYS)]
    band = {}
    for sid in group:
        band.setdefault(sid, len(band) % 2)
    fills = [None, PatternFill("solid", fgColor="F3F6FD")]
    table(ws, 4, cols, rows,
          totals={"Qty (kg)": "sum", "Qty (MT)": "sum", "Cost value ₹": "sum", "Sale value ₹": "sum", "Margin ₹": "sum",
                  "Margin %": lambda a, z: "=IFERROR(T%d/L%d,0)" % (z + 1, z + 1)},
          group_fill=lambda i: fills[band[group[i]]])
    if not rows:
        ws["A5"] = "No sales in this period."
        ws["A5"].font = Font(italic=True, color=MUTED)

    # ================================================================ Sales
    ws = wb.create_sheet("Sales")
    title(ws, "Sales", scope_text)
    by_sale = defaultdict(list)
    for f in flows:
        by_sale[f["sale_deal_id"]].append(f)
    rows = []
    for s in sales:
        fs = by_sale[s["id"]]
        cost = sum(_money(f["qty_g"], f["cost_paise"]) for f in fs)
        value = _money(s["qty_g"], s["rate_paise"])
        sources = "; ".join("%s %s MT @ ₹%.2f (%s)" % (buy_by_id[f["buy_deal_id"]]["party_name"], _fmt_mt(f["qty_g"]),
                                                      f["cost_paise"] / 100, buy_by_id[f["buy_deal_id"]]["ref"]) for f in fs)
        margin = round(value - cost, 2)
        rows.append([
            _d(s["deal_date"]), s["ref"], s["party_name"], s["party_gstin"] or "", s["product"],
            s["warehouse"] or "", _kg(s["qty_g"]), _mt(s["qty_g"]), _r(s["rate_paise"]),
            "GST extra" if s["plus_gst"] else "incl. GST", value,
            round(cost / _kg(s["qty_g"]), 2) if s["qty_g"] else None, cost, margin,
            round(margin / _kg(s["qty_g"]), 2) if s["qty_g"] else None, (margin / cost) if cost else None,
            sources, s["payment_terms"] or "", _d(s["payment_due"]), s["transporter_name"] or "",
            s["delivery_by"] or "", s["freight_by"] or "", s["ex_place"] or "", s["eway"] or "", s["remarks"] or "",
        ])
    cols = [("Date", 12, DATE), ("Sauda No.", 16, None), ("Buyer", 28, None), ("Buyer GSTIN", 17, None),
            ("Product", 30, None), ("Dispatched from", 13, None), ("Qty (kg)", 12, KG), ("Qty (MT)", 10, MT),
            ("Rate ₹/kg", 11, RATE), ("GST", 10, None), ("Sale value ₹", 14, MONEY),
            ("Avg cost ₹/kg", 12, RATE), ("Cost value ₹", 14, MONEY), ("Margin ₹", 14, MONEY),
            ("Margin ₹/kg", 12, MARGIN_RATE), ("Margin %", 10, PCT), ("Material came from", 55, None),
            ("Payment terms", 13, None), ("Payment due", 12, DATE), ("Transporter", 18, None),
            ("Transport arranged by", 12, None), ("Freight paid by", 11, None), ("Ex-Place", 12, None),
            ("E-way bill", 14, None), ("Note", 24, None)]
    table(ws, 4, cols, rows, totals={"Qty (kg)": "sum", "Qty (MT)": "sum", "Sale value ₹": "sum",
                                     "Cost value ₹": "sum", "Margin ₹": "sum",
                                     "Margin %": lambda a, z: "=IFERROR(N%d/M%d,0)" % (z + 1, z + 1)})

    # ================================================================ Purchases
    ws = wb.create_sheet("Purchases")
    title(ws, "Purchases", scope_text + "   ·   includes earlier purchases whose material was sold in this period")
    rows = []
    for b in buys:
        o = outcomes[b["id"]]
        in_window = (not date_from or b["deal_date"] >= date_from) and (not date_to or b["deal_date"] <= date_to)
        left_where = ", ".join("%s %s MT" % (w, _fmt_mt(g)) for w, g in sorted(o["left_where"].items()))
        rows.append([
            _d(b["deal_date"]), b["ref"], "In period" if in_window else "Earlier", b["party_name"], b["party_gstin"] or "",
            b["product"], b["warehouse"] or "", _kg(b["qty_g"]), _mt(b["qty_g"]), _r(b["rate_paise"]),
            "GST extra" if b["plus_gst"] else "incl. GST", _money(b["qty_g"], b["rate_paise"]),
            _kg(o["sold_g"]), "; ".join(o["to"]), round(o["margin"], 2),
            _kg(o["moved_g"]), _kg(o["lost_g"]), _kg(o["found_g"]), _kg(o["left_g"]), left_where,
            _money(o["left_g"], b["rate_paise"]),
            (o["sold_g"] / b["qty_g"]) if b["qty_g"] else None,
            b["payment_terms"] or "", _d(b["payment_due"]), b["transporter_name"] or "", b["remarks"] or "",
        ])
    cols = [("Date", 12, DATE), ("Sauda No.", 16, None), ("Bought", 10, None), ("Supplier", 28, None),
            ("Supplier GSTIN", 17, None), ("Product", 30, None), ("Received into", 12, None),
            ("Qty (kg)", 12, KG), ("Qty (MT)", 10, MT), ("Rate ₹/kg", 11, RATE), ("GST", 10, None),
            ("Purchase value ₹", 15, MONEY), ("Sold (kg)", 12, KG), ("Went to", 55, None), ("Margin earned ₹", 14, MONEY),
            ("Transferred (kg)", 12, KG), ("Written off (kg)", 12, KG), ("Found (kg)", 10, KG),
            ("Still in stock (kg)", 13, KG), ("Where it sits now", 22, None), ("Stock value at cost ₹", 15, MONEY),
            ("% sold", 9, PCT), ("Payment terms", 13, None), ("Payment due", 12, DATE),
            ("Transporter", 18, None), ("Note", 24, None)]
    table(ws, 4, cols, rows, totals={"Qty (kg)": "sum", "Qty (MT)": "sum", "Purchase value ₹": "sum", "Sold (kg)": "sum",
                                     "Margin earned ₹": "sum", "Transferred (kg)": "sum", "Written off (kg)": "sum",
                                     "Found (kg)": "sum", "Still in stock (kg)": "sum", "Stock value at cost ₹": "sum"})

    # ================================================================ Stock Movements
    ws = wb.create_sheet("Stock Movements")
    title(ws, "Transfers and adjustments", scope_text)
    label = {"transfer_out": "Transfer out", "transfer_in": "Transfer in", "loss": "Written off", "gain": "Found"}
    rows = [[_d(m["date"]), label[m["kind"]], m["product"], m["warehouse"],
             (m["counterparty"] or "") if m["kind"].startswith("transfer") else "",
             _kg(m["qty_g"]) if m["qty_g"] > 0 else None, _kg(-m["qty_g"]) if m["qty_g"] < 0 else None,
             m["buy_ref"], m["supplier"], _r(m["rate_paise"]), m["note"] or ""] for m in moves]
    cols = [("Date", 12, DATE), ("Movement", 13, None), ("Product", 30, None), ("Warehouse", 13, None),
            ("From / to warehouse", 16, None), ("In (kg)", 11, KG), ("Out (kg)", 11, KG),
            ("From purchase", 16, None), ("Supplier", 26, None), ("Cost ₹/kg", 11, RATE), ("Reason / note", 30, None)]
    table(ws, 4, cols, rows, totals={"In (kg)": "sum", "Out (kg)": "sum"})
    if not rows:
        ws["A5"] = "No transfers or adjustments in this period."
        ws["A5"].font = Font(italic=True, color=MUTED)

    # ================================================================ Summary (first tab)
    ws = wb.create_sheet("Summary", 0)
    title(ws, "Flow of Material report", scope_text)
    ws["A3"] = "Generated %s" % datetime.now().strftime("%d-%b-%Y %H:%M")
    ws["A3"].font = Font(size=9.5, color=MUTED)
    in_window = [b for b in buys if (not date_from or b["deal_date"] >= date_from) and (not date_to or b["deal_date"] <= date_to)]
    sold_g = sum(s["qty_g"] for s in sales)
    sale_value = sum(_money(s["qty_g"], s["rate_paise"]) for s in sales)
    cost_value = sum(_money(f["qty_g"], f["cost_paise"]) for f in flows)
    bought_g = sum(b["qty_g"] for b in in_window)
    bought_value = sum(_money(b["qty_g"], b["rate_paise"]) for b in in_window)
    kpis = [
        ("Purchases in period", len(in_window), '0'), ("Bought (MT)", _mt(bought_g), MT), ("Purchase value", bought_value, MONEY),
        ("Sales in period", len(sales), '0'), ("Sold (MT)", _mt(sold_g), MT), ("Sale value", sale_value, MONEY),
        ("Cost of material sold", cost_value, MONEY), ("Margin", round(sale_value - cost_value, 2), MONEY),
        ("Margin %", ((sale_value - cost_value) / cost_value) if cost_value else 0, PCT),
        ("Avg margin ₹/kg", round((sale_value - cost_value) / _kg(sold_g), 2) if sold_g else 0, MARGIN_RATE),
        ("Suppliers", len({b["party_id"] for b in in_window}), '0'), ("Buyers", len({s["party_id"] for s in sales}), '0'),
    ]
    r0 = 5
    for i, (name, value, fmt) in enumerate(kpis):
        row, col = r0 + (i // 3) * 2, 1 + (i % 3) * 3
        ws.cell(row=row, column=col, value=name).font = Font(size=9.5, color=MUTED, bold=True)
        cell = ws.cell(row=row + 1, column=col, value=value)
        cell.number_format, cell.font = fmt, Font(size=14, bold=True, color=INK)
    for c in range(1, 10):
        ws.column_dimensions[get_column_letter(c)].width = 16

    # by product
    per = defaultdict(lambda: {"bought_g": 0, "bought_v": 0.0, "sold_g": 0, "sold_v": 0.0, "cost_v": 0.0,
                               "suppliers": set(), "buyers": set()})
    for b in in_window:
        p = per[b["product"]]
        p["bought_g"] += b["qty_g"]; p["bought_v"] += _money(b["qty_g"], b["rate_paise"]); p["suppliers"].add(b["party_name"])
    for s in sales:
        p = per[s["product"]]
        p["sold_g"] += s["qty_g"]; p["sold_v"] += _money(s["qty_g"], s["rate_paise"]); p["buyers"].add(s["party_name"])
    for f in flows:
        per[sale_by_id[f["sale_deal_id"]]["product"]]["cost_v"] += _money(f["qty_g"], f["cost_paise"])
    stock_now = {r["display"]: r["g"] for r in db.q(
        """SELECT s.display, SUM(l.qty_g - l.qty_allocated_g - l.qty_out_g) AS g FROM lots l
           JOIN v_products s ON s.id = l.product_id WHERE l.status = 'open' GROUP BY s.display""")}
    rows = []
    for name in sorted(per):
        p = per[name]
        margin = round(p["sold_v"] - p["cost_v"], 2)
        rows.append([name, _mt(p["bought_g"]), round(p["bought_v"] / _kg(p["bought_g"]), 2) if p["bought_g"] else None,
                     round(p["bought_v"], 2), _mt(p["sold_g"]),
                     round(p["sold_v"] / _kg(p["sold_g"]), 2) if p["sold_g"] else None, round(p["sold_v"], 2),
                     round(p["cost_v"], 2), margin, (margin / p["cost_v"]) if p["cost_v"] else None,
                     _mt(stock_now.get(name, 0)), ", ".join(sorted(p["suppliers"])), ", ".join(sorted(p["buyers"]))])
    top = r0 + 9
    ws.cell(row=top - 1, column=1, value="By product").font = Font(bold=True, size=12, color=INK)
    cols = [("Product", 32, None), ("Bought (MT)", 11, MT), ("Avg buy ₹/kg", 12, RATE), ("Purchase value ₹", 15, MONEY),
            ("Sold (MT)", 11, MT), ("Avg sale ₹/kg", 12, RATE), ("Sale value ₹", 15, MONEY),
            ("Cost of sold ₹", 15, MONEY), ("Margin ₹", 14, MONEY), ("Margin %", 10, PCT),
            ("In stock now (MT)", 13, MT), ("Bought from", 40, None), ("Sold to", 40, None)]
    after = table(ws, top, cols, rows, totals={"Bought (MT)": "sum", "Purchase value ₹": "sum", "Sold (MT)": "sum",
                                               "Sale value ₹": "sum", "Cost of sold ₹": "sum", "Margin ₹": "sum",
                                               "In stock now (MT)": "sum"})
    ws.freeze_panes = None
    # the headline figures sit in columns A, D and G at a large size; make sure
    # the table below never leaves those columns too narrow to show them
    for col in ("A", "D", "G"):
        ws.column_dimensions[col].width = max(ws.column_dimensions[col].width or 0, 20)
    notes = [
        "How to read this report",
        "• Material Flow: each row is one quantity that left a purchase and went into a sale — the supplier and cost on the left, the buyer and sale rate on the right, the margin at the end. Rows of the same sale share a shade.",
        "• A sale made in this period is included with every purchase it drew on, even a purchase made before the period.",
        "• Purchases lists what was bought in the period plus those earlier purchases, with where each one went and what is still in stock.",
        "• Rates are ₹ per kg, basic rate. Margin = (sale rate − cost rate) × quantity. Values are in rupees.",
        "• Filters and totals: every sheet has filters on its header row; the TOTAL row recalculates for the filtered rows.",
    ]
    for i, text in enumerate(notes):
        c = ws.cell(row=after + 2 + i, column=1, value=text)
        c.font = Font(bold=(i == 0), size=10 if i else 11, color=INK if i == 0 else MUTED)

    for sheet in wb.worksheets:
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        if sheet.title != "Summary":
            sheet.print_title_rows = "4:4"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def filename(date_from: Optional[str], date_to: Optional[str]) -> str:
    a = (date_from or "start").replace("-", "")
    b = (date_to or date.today().isoformat()).replace("-", "")
    return "Labdhi-Flow-of-Material_%s-%s.xlsx" % (a, b)
